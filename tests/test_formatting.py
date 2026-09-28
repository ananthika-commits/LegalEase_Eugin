"""
Unit tests for document formatting utilities:
- sanitize_text
- format_docx
- format_pdf
- format_html_preview
"""

import io
import os
import pytest
from docx import Document
from utils.document_formatter import (
    sanitize_text,
    format_docx,
    format_pdf,
    format_html_preview,
)

SAMPLE_DOC_TYPE = "Freelance Work Contract"
SAMPLE_PARTIES = "Jane Doe (Service Provider), TechNova Inc. (Client)"
SAMPLE_DATES = "April 10, 2025"
SAMPLE_TERMS = [
    "Payment of $5,000 upon milestone delivery within 30 days of invoice",
    "Service provider retains preliminary portfolio display rights",
    "Confidentiality must be maintained at all times",
    "Either party may terminate with 15 days notice",
]
SAMPLE_TEXT = """# FREELANCE WORK CONTRACT

**EFFECTIVE DATE:** April 10, 2025

This Agreement is entered into by Jane Doe and TechNova Inc.

## 1. RECITALS
WHEREAS, the Client wishes to engage the Provider for software development;

## 2. OPERATIVE TERMS
- Payment of $5,000 upon milestone delivery within 30 days of invoice.
- Service provider retains preliminary portfolio display rights.
- Confidentiality must be maintained at all times.
- Either party may terminate with 15 days notice.

## 3. GOVERNING LAW
This contract is governed by the laws of California.

## 4. SIGNATURES
Authorized Signature: ____________________
"""


def test_sanitize_text():
    """Verify smart quotes, dashes, bullets, and unicode replacements."""
    dirty_text = "‘Hello’ “World” — en-dash – bullet • non-breaking\u00a0space"
    clean = sanitize_text(dirty_text)
    assert "'" in clean
    assert '"' in clean
    assert "-" in clean
    assert "*" in clean
    # Check all characters are latin-1 safe
    for ch in clean:
        assert ord(ch) < 256


def test_format_docx_generation():
    """Verify docx generation produces a valid, readable .docx document."""
    logo_path = os.path.join(os.path.dirname(__file__), "..", "assets", "logo.png")
    buf = format_docx(
        text=SAMPLE_TEXT,
        doc_type=SAMPLE_DOC_TYPE,
        terms=SAMPLE_TERMS,
        parties=SAMPLE_PARTIES,
        effective_date=SAMPLE_DATES,
        logo_path=logo_path if os.path.exists(logo_path) else None
    )
    assert isinstance(buf, io.BytesIO)
    content = buf.getvalue()
    assert len(content) > 1000  # Non-trivial docx file size

    # Verify python-docx can open and read the generated document
    doc = Document(buf)
    paragraphs = [p.text for p in doc.paragraphs if p.text]
    assert any(SAMPLE_DOC_TYPE.upper() in p for p in paragraphs)
    
    # Check that tables exist (metadata table and terms table)
    assert len(doc.tables) >= 1


def test_format_pdf_generation():
    """Verify PDF generation outputs valid PDF bytes with standard header."""
    logo_path = os.path.join(os.path.dirname(__file__), "..", "assets", "logo.png")
    buf = format_pdf(
        text=SAMPLE_TEXT,
        doc_type=SAMPLE_DOC_TYPE,
        terms=SAMPLE_TERMS,
        parties=SAMPLE_PARTIES,
        effective_date=SAMPLE_DATES,
        logo_path=logo_path if os.path.exists(logo_path) else None
    )
    assert isinstance(buf, io.BytesIO)
    content = buf.getvalue()
    assert len(content) > 500
    # PDF files start with %PDF
    assert content.startswith(b"%PDF")


def test_format_html_preview():
    """Verify HTML preview returns dark card container with stylized elements."""
    html = format_html_preview(SAMPLE_TEXT, SAMPLE_DOC_TYPE)
    assert "LegalEase Verified Draft" in html
    assert "FREELANCE WORK CONTRACT" in html
    assert "<h1" in html
    assert "<h2" in html
    assert "RECITALS" in html
